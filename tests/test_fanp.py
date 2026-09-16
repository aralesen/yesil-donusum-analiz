# -*- coding: utf-8 -*-
"""
FANP motoru testleri. Çalıştırma:  python -m pytest -q tests

Testler tek bir veri setine bağlı değildir: her testte küme, kriter ve strateji sayısı, puan ölçeği,
bulanıklık genişliği ve iç bağımlılık rastgele seçilir.
"""
import io
import os
import sys

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
import fanp_motor as m  # noqa: E402

SCALES = [(1, 5, 1), (1, 7, 1), (1, 9, 1), (1, 9, 2), (1, 10, 0.5), (2, 9, 1), (1, 9, 0)]


# ----------------------------------------------------------------------------- yardımcılar
def random_model(rng, dependence=None):
    K = int(rng.integers(1, 7))
    A = int(rng.integers(2, 8))
    lo, hi, sp = SCALES[rng.integers(len(SCALES))]
    clusters = [(f"K{k + 1}", f"Küme {k + 1}", f"Ana {k + 1}") for k in range(K)]
    criteria = []
    for k in range(K):
        for j in range(int(rng.integers(1, 8))):
            criteria.append((f"K{k + 1}.{j + 1}", f"Soru {k + 1}-{j + 1}", f"Kriter {k + 1}.{j + 1}", f"K{k + 1}"))
    alts = [(f"S{a + 1}", f"Strateji {a + 1}") for a in range(A)]
    mat = rng.integers(lo, hi + 1, (len(criteria), A)).astype(float)
    if dependence is None:
        dependence = rng.random() < 0.4
    dep = rng.integers(0, 10, (K, K)).astype(float) + np.eye(K) if dependence else None
    return m.Model(clusters, criteria, alts, mat, lo, hi, sp, dep, 'Test')


def random_scores(rng, model, D):
    lo, hi = model.scale_min, model.scale_max
    R = rng.uniform(lo, hi, (D, len(model.codes)))
    C = rng.uniform(lo, hi, (D, len(model.cl_codes)))
    R[: D // 2] = np.round(R[: D // 2])                     # yarısı tam sayı
    C[: D // 2] = np.round(C[: D // 2])
    return np.clip(R, lo, hi), np.clip(C, lo, hi)


def reference_fanp(model, r, c, sentez):
    """Döngülerle yazılmış bağımsız uygulama: bulanık AHP adımları + tam süpermatris kuvveti."""
    lo, hi, sp = model.scale_min, model.scale_max, model.spread
    tfn = lambda x: (max(lo, x - sp), x, min(hi, x + sp))

    def priority(vals):
        T = [tfn(x) for x in vals]
        n = len(T)
        mats = [[[1.0 if i == j else T[i][0] / T[j][2] for j in range(n)] for i in range(n)],
                [[1.0 if i == j else T[i][1] / T[j][1] for j in range(n)] for i in range(n)],
                [[1.0 if i == j else T[i][2] / T[j][0] for j in range(n)] for i in range(n)]]
        cs = [[sum(M[i][j] for i in range(n)) for j in range(n)] for M in mats]
        pl = [sum(mats[0][i][j] / cs[2][j] for j in range(n)) / n for i in range(n)]
        pm = [sum(mats[1][i][j] / cs[1][j] for j in range(n)) / n for i in range(n)]
        pu = [sum(mats[2][i][j] / cs[0][j] for j in range(n)) / n for i in range(n)]
        return [pl, pm, pu]

    n, K, A = len(model.codes), len(model.cl_codes), len(model.alt_codes)
    cl_of = [model.cl_codes.index(cr[3]) for cr in model.criteria]
    loc = [[0.0] * n for _ in range(3)]
    for k in range(K):
        idx = [i for i in range(n) if cl_of[i] == k]
        p = priority([r[i] for i in idx])
        for t in range(3):
            for q, i in enumerate(idx):
                loc[t][i] = p[t][q]
    clw = priority(list(c))
    altn = [[[0.0] * A for _ in range(n)] for _ in range(3)]
    for i in range(n):
        T = [tfn(x) for x in model.alt_matrix[i]]
        sl, sm, su = sum(t[0] for t in T), sum(t[1] for t in T), sum(t[2] for t in T)
        for a in range(A):
            altn[0][i][a], altn[1][i][a], altn[2][i][a] = T[a][0] / su, T[a][1] / sm, T[a][2] / sl
    dz = lambda a, b, c_: (a + 2 * b + c_) / 4
    if sentez == 'durulastirilmis':
        cr = [dz(loc[0][i], loc[1][i], loc[2][i]) for i in range(n)]
        for k in range(K):
            s = sum(cr[i] for i in range(n) if cl_of[i] == k)
            for i in range(n):
                if cl_of[i] == k:
                    cr[i] /= s
        cc = [dz(clw[0][k], clw[1][k], clw[2][k]) for k in range(K)]
        cc = [x / sum(cc) for x in cc]
        an = []
        for i in range(n):
            row = [dz(altn[0][i][a], altn[1][i][a], altn[2][i][a]) for a in range(A)]
            an.append([x / sum(row) for x in row])
        loc, clw, altn, comps = [cr] * 3, [cc] * 3, [an] * 3, 1
    else:
        comps = 3
    out = []
    for t in range(comps):
        N = 1 + n + A
        W = np.zeros((N, N))
        for i in range(n):
            W[1 + i, 0] = clw[t][cl_of[i]] * loc[t][i]
        alpha = m.BAGIMLILIK_AGIRLIGI if model.dependence is not None else 0.0
        if model.dependence is not None:
            Dm = model.dependence
            ln = list(loc[t])
            for k in range(K):
                s = sum(ln[i] for i in range(n) if cl_of[i] == k)
                for i in range(n):
                    if cl_of[i] == k:
                        ln[i] /= s
            for i in range(n):
                for j in range(n):
                    W[1 + i, 1 + j] = alpha * Dm[cl_of[i], cl_of[j]] / Dm[:, cl_of[j]].sum() * ln[i]
        for j in range(n):
            for a in range(A):
                W[1 + n + a, 1 + j] = (1 - alpha) * altn[t][j][a]
        for a in range(A):
            W[1 + n + a, 1 + n + a] = 1.0
        L = W.copy()
        for _ in range(60):                                   # W^(2^60): kesin yakınsama
            L2 = L @ L
            if np.max(np.abs(L2 - L)) < 1e-16:
                break
            L = L2
        out.append(L[1 + n:, 0])
    return out[0] if comps == 1 else dz(*out)


def write_survey_excel(rng, model, R, C, ids, layout_seed):
    """Aynı veriyi farklı düzenlerle Excel'e yazar: sütun sırası karışık, birden çok ID bloğu,
    başlık üstünde boş satırlar, fazladan sütunlar, virgüllü ondalık metinler."""
    lr = np.random.default_rng(layout_seed)
    items = [(code, model.header_of[code], R[:, i]) for i, code in enumerate(model.codes)] + \
            [(code, model.header_of[code], C[:, k]) for k, code in enumerate(model.cl_codes)]
    order = lr.permutation(len(items))
    n_blocks = int(lr.integers(1, 4))
    splits = np.array_split(order, n_blocks)
    offset = int(lr.integers(0, 4))
    header, cols = [], []
    for b, part in enumerate(splits):
        perm = lr.permutation(len(ids))                      # her blokta satır sırası farklı
        header.append('ID')
        cols.append([ids[p] for p in perm])
        for idx in part:
            code, h, vals = items[idx]
            header.append(h if lr.random() < 0.7 else f"  {h.upper()} ")
            col = []
            for p in perm:
                v = vals[p]
                col.append(str(v).replace('.', ',') if lr.random() < 0.1 else v)
            cols.append(col)
        header.append('Ortalama' if lr.random() < 0.5 else None)
        cols.append([lr.random() for _ in ids])
    rows = [[None] * len(header) for _ in range(offset)] + [header] + \
           [[cols[j][i] for j in range(len(header))] for i in range(len(ids))]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as xw:
        pd.DataFrame([['Company ID', 'Company Size', 'Sector', 'Motivation']] +
                     [[i, 'Small', 'Metal', 'Cost reduction'] for i in ids]).to_excel(xw, sheet_name='Demografi', index=False, header=False)
        pd.DataFrame(rows).to_excel(xw, sheet_name='Puanlar', index=False, header=False)
    return buf.getvalue()


# ----------------------------------------------------------------------------- testler
@pytest.mark.parametrize('seed', range(40))
@pytest.mark.parametrize('sentez', ['durulastirilmis', 'bulanik'])
def test_bagimsiz_referansla_ayni(seed, sentez):
    rng = np.random.default_rng(seed)
    model = random_model(rng)
    R, C = random_scores(rng, model, 6)
    got = m.run_fanp(model, R, C, sentez)['net']
    full = m.run_fanp(model, R, C, sentez, full_limit=True)['net']
    for d in range(len(R)):
        ref = reference_fanp(model, R[d], C[d], sentez)
        assert np.allclose(got[d], ref, atol=1e-10), (seed, d)
        assert np.allclose(full[d], ref, atol=1e-10), (seed, d)


@pytest.mark.parametrize('seed', range(30))
def test_olasilik_ve_toplam_ozellikleri(seed):
    rng = np.random.default_rng(100 + seed)
    model = random_model(rng)
    R, C = random_scores(rng, model, 20)
    res = m.run_fanp(model, R, C, 'durulastirilmis')
    assert np.all(res['net'] > 0)
    assert np.allclose(res['cluster'].sum(1), 1)
    assert np.allclose(res['global'].sum(1), 1)
    if model.dependence is None:
        assert np.allclose(res['net'].sum(1), 1)
    for k, cols in enumerate(model.members):                 # kümenin global ağırlık toplamı = küme ağırlığı
        assert np.allclose(res['global'][:, cols].sum(1), res['cluster'][:, k])


@pytest.mark.parametrize('seed', range(25))
def test_sira_degismezligi(seed):
    """Kriterlerin, kümelerin ve stratejilerin sırası sonucu değiştirmemeli."""
    rng = np.random.default_rng(200 + seed)
    model = random_model(rng)
    R, C = random_scores(rng, model, 5)
    base = m.run_fanp(model, R, C)
    pc = rng.permutation(len(model.codes))
    pk = rng.permutation(len(model.cl_codes))
    pa = rng.permutation(len(model.alt_codes))
    dep = None if model.dependence is None else model.dependence[np.ix_(pk, pk)]
    m2 = m.Model([model.clusters[k] for k in pk], [model.criteria[i] for i in pc], [model.alternatives[a] for a in pa],
                 model.alt_matrix[np.ix_(pc, pa)], model.scale_min, model.scale_max, model.spread, dep)
    res = m.run_fanp(m2, R[:, pc], C[:, pk])
    assert np.allclose(res['net'], base['net'][:, pa], atol=1e-12)
    assert np.allclose(res['global'], base['global'][:, pc], atol=1e-12)


@pytest.mark.parametrize('seed', range(25))
def test_baskin_strateji_geride_kalmaz(seed):
    """Bir strateji her kriterde diğerinden en az onun kadar puan alıyorsa puanı düşük çıkamaz."""
    rng = np.random.default_rng(300 + seed)
    model = random_model(rng, dependence=False)
    mat = model.alt_matrix.copy()
    mat[:, 0] = np.maximum(mat[:, 0], mat[:, 1])
    model = m.Model(model.clusters, model.criteria, model.alternatives, mat, model.scale_min, model.scale_max, model.spread)
    R, C = random_scores(rng, model, 10)
    for sentez in ('durulastirilmis', 'bulanik'):
        net = m.run_fanp(model, R, C, sentez)['net']
        assert np.all(net[:, 0] >= net[:, 1] - 1e-12)


@pytest.mark.parametrize('seed', range(15))
def test_kume_puani_artinca_agirligi_azalmaz(seed):
    rng = np.random.default_rng(400 + seed)
    model = random_model(rng, dependence=False)
    if len(model.cl_codes) < 2:
        return
    R, C = random_scores(rng, model, 1)
    k = int(rng.integers(len(model.cl_codes)))
    prev = -1
    for v in np.linspace(model.scale_min, model.scale_max, 9):
        C2 = C.copy(); C2[0, k] = v
        w = m.run_fanp(model, R, C2)['cluster'][0, k]
        assert w >= prev - 1e-12
        prev = w


def test_esit_puan_esit_agirlik_ve_tek_elemanli_kume():
    model = m.Model([('X', 'X', 'X'), ('Y', 'Y', 'Y')], [('X.1', 'a', 'a', 'X'), ('Y.1', 'b', 'b', 'Y'), ('Y.2', 'c', 'c', 'Y')],
                    [('P', 'P'), ('Q', 'Q')], [[3, 3], [3, 3], [3, 3]], 1, 5, 1)
    res = m.run_fanp(model, [4, 2, 2], [3, 3], 'durulastirilmis')
    assert np.allclose(res['global'][0], [0.5, 0.25, 0.25])
    assert np.allclose(res['net'][0], [0.5, 0.5])


@pytest.mark.parametrize('seed', range(10))
def test_saglamlik(seed):
    rng = np.random.default_rng(500 + seed)
    model = random_model(rng)
    R, C = random_scores(rng, model, 1)
    p = m.robustness(model, R[0], C[0], 300, seed)
    assert p.shape == (len(model.alt_codes),) and np.isclose(p.sum(), 1)
    sharp = model.with_scale(model.scale_min, model.scale_max, 0)
    p0 = m.robustness(sharp, R[0], C[0], 50, seed)
    assert p0[m.run_fanp(sharp, R, C)['net'][0].argmax()] == 1.0


@pytest.mark.parametrize('seed', range(20))
def test_excel_okuma_duzenden_bagimsiz(seed):
    rng = np.random.default_rng(600 + seed)
    model = random_model(rng)
    R, C = random_scores(rng, model, 7)
    ids = [str(i) for i in rng.choice(900, 7, replace=False) + 1]
    content = write_survey_excel(rng, model, R, C, ids, seed)
    got_model, survey, demo, rep, info = m.read_workbook(content, model)
    assert set(survey.index) == set(ids)
    for d, fid in enumerate(ids):
        assert np.allclose(survey.loc[fid, model.codes].values.astype(float), R[d])
        assert np.allclose(survey.loc[fid, model.cl_codes].values.astype(float), C[d])
    res = m.analyze(got_model, survey, demo, rep, simulations=0)
    assert len(res['firms']) == len(ids)
    assert np.allclose(res['firms'][[f'{a} payı (%)' for a in model.alt_codes]].sum(axis=1), 100)


def test_eksik_ve_aralik_disi_puanlar_raporlanir():
    model = m.default_model()
    rng = np.random.default_rng(3)
    R, C = random_scores(rng, model, 5)
    R[1, 3] = np.nan
    R[2, 0] = model.scale_max + 3
    survey = pd.DataFrame(np.hstack([R, C]), columns=model.needed, index=['1', '2', '3', '4', '5'])
    res = m.analyze(model, survey, pd.DataFrame(columns=['Scale', 'Sector', 'Motivation']), simulations=0)
    assert list(res['firms']['ID']) == ['1', '4', '5']
    kinds = set(res['report']['Tür'])
    assert {'Eksik puan', 'Aralık dışı'} <= kinds


@pytest.mark.parametrize('bad', ['dup_code', 'orphan', 'shape', 'range', 'one_alt', 'empty_cluster', 'dep_shape'])
def test_tutarsiz_model_reddedilir(bad):
    cl = [('X', 'X', 'hx'), ('Y', 'Y', 'hy')]
    cr = [('X.1', 'a', 'a', 'X'), ('Y.1', 'b', 'b', 'Y')]
    al = [('P', 'P'), ('Q', 'Q')]
    mat = [[3, 4], [5, 6]]
    kw = {}
    if bad == 'dup_code':
        cr = [('X.1', 'a', 'a', 'X'), ('X.1', 'b', 'b', 'Y')]
    elif bad == 'orphan':
        cr = [('X.1', 'a', 'a', 'X'), ('Z.1', 'b', 'b', 'Z')]
    elif bad == 'shape':
        mat = [[3, 4, 5], [5, 6, 7]]
    elif bad == 'range':
        mat = [[3, 40], [5, 6]]
    elif bad == 'one_alt':
        al, mat = [('P', 'P')], [[3], [5]]
    elif bad == 'empty_cluster':
        cl = cl + [('Z', 'Z', 'hz')]
    elif bad == 'dep_shape':
        kw['dependence'] = np.ones((3, 3))
    with pytest.raises(m.ModelError):
        m.Model(cl, cr, al, mat, 1, 9, 1, **kw)


def test_varsayilan_model_gecerlilik_testi():
    vt = m.validity_test(m.default_model())
    assert (vt['Kazanan (durulastirilmis)'] == vt['Beklenen']).all()


@pytest.mark.parametrize('seed', range(10))
def test_tutarlilik_orani(seed):
    """Puanlardan oranla türetilen matrisler tanım gereği tutarlıdır: CR = 0."""
    rng = np.random.default_rng(800 + seed)
    model = random_model(rng)
    R, C = random_scores(rng, model, 8)
    crs = m.consistency_ratios(model, R, C)
    assert crs.shape == (8, len(model.cl_codes) + 1)
    assert np.all(np.abs(crs.values) < 1e-9)


# ----------------------------------------------------------------------------- şablonlar
def _fill_single(model, R, C, info):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(m.single_firm_template_excel(model)))
    ws = wb['Firma']
    hdr = {str(ws.cell(2, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value}
    score_c = next(c for h, c in hdr.items() if h.startswith('Puan'))
    weight_c = next(c for h, c in hdr.items() if h.startswith('Ağırlık'))
    for i in range(len(model.codes)):
        if R[i] is not None and not np.isnan(R[i]):
            ws.cell(3 + i, score_c, float(R[i]))
    for k in range(len(model.cl_codes)):
        ws.cell(3 + k, weight_c, float(C[k]))
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, weight_c - 3).value
        if lab in info:
            ws.cell(r, weight_c - 2, info[lab])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _template_text(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return [str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]


@pytest.mark.parametrize('seed', range(12))
def test_sablonlarda_strateji_yok(seed):
    """Stratejiler sonuçtur: iki şablonda da strateji kodu, adı ya da uzman puanı bulunmaz."""
    rng = np.random.default_rng(1000 + seed)
    model = random_model(rng) if seed else m.default_model()
    for data in (m.single_firm_template_excel(model), m.multi_firm_template_excel(model)):
        cells = _template_text(data)
        for code, name in model.alternatives:
            assert all(c.strip() != code and not c.startswith(code + ':') and name not in c for c in cells)
    cells = set(_template_text(m.single_firm_template_excel(model)))
    assert set(model.codes) <= cells and set(model.cl_codes) <= cells


@pytest.mark.parametrize('seed', range(15))
def test_tek_firma_sablonu(seed):
    rng = np.random.default_rng(1100 + seed)
    model = random_model(rng)
    R, C = random_scores(rng, model, 1)
    info = {'Firma adı': f'Firma {seed}', 'Ölçek': 'Orta', 'Sektör': 'Metal', 'Motivasyon': 'Cost reduction'}
    _, survey, demo, rep, _ = m.read_workbook(_fill_single(model, R[0], C[0], info), model, kind='tek')
    fid = f'Firma {seed}'
    assert list(survey.index) == [fid]
    assert np.allclose(survey.loc[fid, model.codes].values.astype(float), R[0])
    assert np.allclose(survey.loc[fid, model.cl_codes].values.astype(float), C[0])
    res = m.analyze(model, survey, demo, rep, simulations=0)
    row = res['firms'].iloc[0]
    assert row['Ölçek'] == 'Orta' and row['Sektör'] == 'Metal' and row['Motivasyon'] == 'Cost reduction'
    expect = m.run_fanp(model, R, C)['net'][0]
    shares = row[[f'{a} payı (%)' for a in model.alt_codes]].values.astype(float) / 100
    assert np.allclose(shares, expect / expect.sum())


@pytest.mark.parametrize('seed', range(15))
def test_coklu_firma_sablonu(seed):
    import openpyxl
    rng = np.random.default_rng(1200 + seed)
    model = random_model(rng)
    n = int(rng.integers(1, 12))
    R, C = random_scores(rng, model, n)
    wb = openpyxl.load_workbook(io.BytesIO(m.multi_firm_template_excel(model)))
    firms, sc = wb['Firmalar'], wb['Puanlar']
    hdr = {sc.cell(2, c).value: c for c in range(1, sc.max_column + 1) if sc.cell(2, c).value}
    rows = rng.choice(30, n, replace=False)
    for d, r in enumerate(rows):
        firms.cell(2 + int(r), 2, ['Mikro', 'Küçük', 'Orta', 'Büyük'][d % 4])
        for j, code in enumerate(model.codes):
            sc.cell(3 + int(r), hdr[model.header_of[code]], float(R[d, j]))
        for k, code in enumerate(model.cl_codes):
            sc.cell(3 + int(r), hdr[model.header_of[code]], float(C[d, k]))
    buf = io.BytesIO()
    wb.save(buf)
    _, survey, demo, rep, _ = m.read_workbook(buf.getvalue(), model, kind='coklu')
    assert sorted(survey.index, key=m.id_key) == sorted([str(int(r) + 1) for r in rows], key=m.id_key)
    for d, r in enumerate(rows):
        assert np.allclose(survey.loc[str(int(r) + 1), model.codes].values.astype(float), R[d])
        assert np.allclose(survey.loc[str(int(r) + 1), model.cl_codes].values.astype(float), C[d])
    assert not [x for x in rep if x['Tür'] in ('Eksik puan', 'Veri')]
    assert len(m.analyze(model, survey, demo, rep, simulations=0)['firms']) == n


def test_calisma_dosyasi_duzeni_tek_firma():
    """Elle kurulmuş çalışma dosyası düzeni: bloklar arasında boşluk yok, İngilizce başlıklar, etiketli kodlar ve
    sağda strateji sütunları. Sadece kriter puanları ve ana başlık ağırlıkları okunur, strateji sütunları yok sayılır."""
    model = m.default_model()
    rng = np.random.default_rng(7)
    R = rng.integers(1, 10, len(model.codes))
    C = rng.integers(1, 10, len(model.cl_codes))
    rows = [['Kriter Kodu', 'Kriter Adı', 'Puan(1-9)', None, 'Cluster_Code', 'Cluster_Weight_Score', 'Kriter Kodu',
             'A1 (Teknoloji)', 'A2 (Döngüsel)', 'A3 (Sosyal)', 'A4 (Yasal)']]
    for i, code in enumerate(model.codes):
        row = [code, f"ad {i}", int(R[i]), None, None, None, f"{code} (etiket)", 1, 2, 3, 4]
        if i < len(model.cl_codes):
            row[4], row[5] = model.cl_codes[i], int(C[i])
        rows.append(row)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as xw:
        pd.DataFrame(rows).to_excel(xw, sheet_name='MAIN_DATA', index=False, header=False)
    _, survey, _, rep, _ = m.read_workbook(buf.getvalue(), model, kind='tek')
    assert np.allclose(survey.iloc[0][model.codes].values.astype(float), R)
    assert np.allclose(survey.iloc[0][model.cl_codes].values.astype(float), C)
    assert not rep


def test_tek_firma_eksik_puan_raporlanir():
    model = m.default_model()
    rng = np.random.default_rng(9)
    R, C = random_scores(rng, model, 1)
    R[0, 4] = np.nan
    _, survey, demo, rep, _ = m.read_workbook(_fill_single(model, R[0], C[0], {}), model, kind='tek')
    res = m.analyze(model, survey, demo, rep, simulations=0)
    assert res['firms'].empty and 'Eksik puan' in set(res['report']['Tür'])


def test_yanlis_tur_dosyasi_acik_hata():
    model = m.default_model()
    with pytest.raises(ValueError):
        m.read_workbook(m.multi_firm_template_excel(model), model, kind='tek')
